"""生成 Englearn 使用教程（DOCX），随后交给 LibreOffice Kit 转 PDF。

图片全部来自 docs/images/，由 scripts/shoot*.py 用无头 Chrome 实拍生成。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "docs" / "images"
OUT = ROOT / "docs" / "Englearn-使用教程.docx"

GREEN = "0B7F63"
DARK = "1B2421"
GRAY = "6B7A75"
FONT = "微软雅黑"
MONO = "Consolas"
VERSION = "v0.1.0"
DOC_DATE = "2026-10-25"

_fig_no = {"n": 0}


# ---------------------------------------------------------------- 样式工具
def set_run(run, size=10.5, bold=False, color=None, name=FONT, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def style_setup(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    for name, size, color, before, after in (
        ("Heading 1", 16, GREEN, 20, 10),
        ("Heading 2", 12.5, DARK, 14, 6),
    ):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:ascii"), FONT)
        rfonts.set(qn("w:hAnsi"), FONT)
        rfonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def shade_para(p, fill: str) -> None:
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def para_border(p, color="DCE7E3", size=6) -> None:
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), color)
        pbdr.append(el)
    ppr.append(pbdr)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def p(doc, text, size=10.5, bold=False, color=None, space_after=6):
    par = doc.add_paragraph()
    run = par.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    par.paragraph_format.space_after = Pt(space_after)
    return par


def bullet(doc, text, bold_prefix=""):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r1 = par.add_run(bold_prefix)
        set_run(r1, bold=True)
    r2 = par.add_run(text)
    set_run(r2)
    return par


def numbered(doc, n, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.first_line_indent = Cm(-0.8)
    pf.space_after = Pt(3)
    r1 = par.add_run(f"{n}.　")
    set_run(r1, bold=True, color=GREEN)
    r2 = par.add_run(text)
    set_run(r2)
    return par


def code(doc, text, size=8.5):
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(6 if i == 0 else 0)
        pf.space_after = Pt(6 if i == len(lines) - 1 else 0)
        pf.left_indent = Cm(0.35)
        pf.line_spacing = 1.15
        shade_para(par, "F4F7F6")
        if i < len(lines) - 1:
            para_border(par, "F4F7F6", 4)
        run = par.add_run(line or " ")
        set_run(run, size=size, name=MONO, color="2B3A35")


def callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "EAF7F2")
    par = cell.paragraphs[0]
    r1 = par.add_run(f"{title}　")
    set_run(r1, size=10, bold=True, color=GREEN)
    r2 = par.add_run(text)
    set_run(r2, size=10, color="27423A")
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(6)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def crop_image(src_name: str, box: tuple, out_name: str) -> str:
    from PIL import Image

    im = Image.open(IMG / src_name)
    im.crop(box).save(IMG / out_name)
    return out_name


def figure(doc, filename: str, caption: str, width: float = 16.2) -> None:
    _fig_no["n"] += 1
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(IMG / filename), width=Cm(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    run = cap.add_run(f"图 {_fig_no['n']}　{caption}")
    set_run(run, size=9, color=GRAY)


def table(doc, headers, rows, widths=None):
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)

    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade_cell(cell, "10A37F")
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(2)
        head_run = par.add_run(head)
        set_run(head_run, size=9, bold=True, color="FFFFFF")

    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            cell_run = par.add_run(str(value))
            set_run(cell_run, size=9)

    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def footer_page_number(doc):
    footer = doc.sections[0].footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("Englearn 使用教程　·　")
    set_run(run, size=8.5, color=GRAY)
    run2 = par.add_run()
    set_run(run2, size=8.5, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2._r.append(begin)
    run2._r.append(instr)
    run2._r.append(end)


# ---------------------------------------------------------------- 正文
def build() -> Document:
    doc = Document()
    style_setup(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    footer_page_number(doc)

    # ---------------- 封面
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(46)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Englearn 使用教程")
    set_run(run, size=30, bold=True, color=DARK)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(2)
    run = sub.add_run("本地英语精读助手　·　划词悬浮窗　·　词汇卡片　·　生词沉淀")
    set_run(run, size=11.5, color=GRAY)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(26)
    run = meta.add_run(f"{VERSION}　｜　适用 Windows 10 / 11　｜　{DOC_DATE}")
    set_run(run, size=9.5, color=GRAY)

    # ---------------- 一
    h1(doc, "一、Englearn 是什么")
    p(doc, "Englearn 是一个只在你电脑上运行的英语精读工具。你在浏览器、PDF 或 Word 里选中一段英文，"
           "光标旁会弹出一个小窗；选「翻译」看释义，选「长难句」看句子结构，选「收藏」存进本地生词本。"
           "所有数据都留在本机，随时可以导出成 Markdown 等格式。")
    bullet(doc, "：服务只监听 127.0.0.1，数据库是一个本地 SQLite 文件。", "纯本地")
    bullet(doc, "：只用 Python 标准库，不需要 pip 安装任何包。", "零依赖")
    bullet(doc, "：任意软件里选中英文，光标旁直接弹出操作窗。", "划词即用")
    bullet(doc, "：Markdown / JSON / CSV / Anki / HTML 五种，数据随时带走。", "可导出")
    figure(doc, "01-main.png", "主界面：左侧输入与结果，右侧生词本 / 长难句 / 历史 / 日历")

    # ---------------- 二
    h1(doc, "二、安装与启动")
    h2(doc, "2.1 三种启动方式")
    table(doc,
          ["方式", "命令", "适用场景", "功能"],
          [["Windows 双击", "start.cmd", "日常使用", "完整（含悬浮窗与热键）"],
           ["Linux / macOS", "./run.sh", "只用网页界面", "仅网页"],
           ["Docker", "docker compose up -d", "服务器 / 隔离环境", "仅网页"]],
          widths=[3.4, 5.0, 4.2, 4.6])
    p(doc, "推荐 Windows 用户直接双击 start.cmd：脚本会自动查找 Python，启动本地服务并打开浏览器。")

    h2(doc, "2.2 控制台输出")
    p(doc, "启动成功时控制台大致如下（端口与热键以实际输出为准）：")
    code(doc, """  [Englearn] interpreter: python
========================================================
  Englearn v0.1.0 已启动
  地址：http://127.0.0.1:8765
  模式：真实模型　模型：deepseek-chat
  关闭本窗口或按 Ctrl+C 即可退出
========================================================
  热键：alt+q（选中英文后按下取词）
  划词：选中英文后鼠标旁会弹出操作窗""")

    h2(doc, "2.3 验证安装")
    p(doc, "不用 API Key 也能验证部署是否正常：")
    code(doc, "python scripts/selftest.py")
    p(doc, "该脚本会在临时目录拉起一个演示服务，跑完 11 项检查后自动关闭，不碰你的真实数据。"
           "全部通过时最后一行显示「11/11 项通过」。")

    # ---------------- 三
    h1(doc, "三、首次配置")
    p(doc, "点右上角「设置」，填入模型服务信息。默认对接 DeepSeek，也兼容任意 OpenAI 协议服务。")
    figure(doc, crop_image("02-settings.png", (688, 240, 1474, 1131), "_crop-settings.png"),
           "设置面板：API Key、Base URL、模型与热键", width=9.8)
    numbered(doc, 1, "API Key：填入服务商密钥；保存后显示占位提示，留空表示不修改。")
    numbered(doc, 2, "Base URL：默认 https://api.deepseek.com；换 OpenAI 则填 https://api.openai.com。")
    numbered(doc, 3, "模型：与 Base URL 配套，DeepSeek 填 deepseek-chat。")
    numbered(doc, 4, "取词热键：默认 alt+d，修改后需重启生效。")
    numbered(doc, 5, "没有 Key 时先勾选「Mock 演示模式」，可用示例数据体验全部界面。")

    # ---------------- 四
    h1(doc, "四、划词取词")
    h2(doc, "4.1 选中即弹窗")
    p(doc, "在任意软件里用鼠标选中一段英文（拖动超过 20 像素，或双击选词），光标旁会弹出操作窗：")
    figure(doc, "08-popup.png", "选中英文后的悬浮操作窗", width=12.6)
    table(doc,
          ["按钮", "效果"],
          [["翻译", "窗内直接显示译文或按词性分组的释义"],
           ["长难句", "窗内显示主干、从句与难点说明"],
           ["收藏", "单词进生词本，句子进长难句库"],
           ["页面", "把完整结果推送到浏览器页面，适合看分层结构"]],
          widths=[3.0, 13.2])
    p(doc, "点击窗外任意位置、右上角 ✕ 或按 Esc 关闭；25 秒无操作也会自动收起。")

    h2(doc, "4.2 全局热键")
    p(doc, "默认 Alt+D。若被其他程序占用，启动时会自动降级到 Alt+Q、Ctrl+Alt+E 等组合，"
           "实际生效的键会显示在控制台与页面输入框上方。")

    h2(doc, "4.3 剪贴板与手动粘贴")
    p(doc, "在设置里勾选「监听剪贴板」后，手动 Ctrl+C 复制的英文也会触发分析；"
           "也可以直接把英文粘贴进输入框，点「分析」或按 Ctrl+Enter。")

    # ---------------- 五
    h1(doc, "五、词汇卡片")
    p(doc, "分析单词或短语时，输出的是七段式卡片：音标（英/美）、释义、熟词生义、短语搭配、"
           "例句、派生词、类似词。没有对应内容的段落显示「暂无」，不会编造。")
    figure(doc, "21-word-card.png", "词汇卡片：七项内容一次给全")

    # ---------------- 六
    h1(doc, "六、长难句解析")
    p(doc, "超过 12 个词或带句末标点的文本按句子处理。结果按「主干 → 从句 → 修饰」分层展示，"
           "每层标注类型与作用，最后给出难点词汇与一句话难点说明。")
    figure(doc, "17-sentence-page.png", "长难句详情：主干、逐层从句、修饰与难点词汇")
    table(doc,
          ["区块", "内容", "怎么用"],
          [["主干", "主语 + 谓语 + 宾语", "先读懂这一句，其余都是附加信息"],
           ["从句", "定语 / 状语 / 宾语 / 表语等", "看标签与说明，理清它修饰谁"],
           ["修饰", "分词短语、插入语、后置定语", "常是「读不懂」的真正原因"],
           ["难点词汇", "句中影响理解的关键词", "点高亮词可直接生成词汇卡片"]],
          widths=[3.4, 6.6, 6.2])

    # ---------------- 七
    h1(doc, "七、生词本与详情")
    p(doc, "右侧「生词」标签页集中管理收藏的单词：同名重复收藏只累加次数，圆点表示掌握状态"
           "（黄=新词、蓝=学习中、绿=已掌握）。点任意一条，左栏整页展示详情。")
    figure(doc, "15-word-page.png", "生词详情：音标、例句、备注、来源与出现次数")
    p(doc, "详情页可以改状态、写备注（回车保存）、删除，也可以点「生成词汇卡片」看七段式解析。"
           "左上角「← 返回」或 Esc 回到输入区。")

    # ---------------- 八
    h1(doc, "八、句子标记")
    p(doc, "在收藏的长难句里，可以选中句中的陌生词并点「标记陌生词」，该词会在句子里高亮显示，"
           "下方生成标签，右侧列表显示标记数量。点标签上的 × 取消标记。")
    figure(doc, "20-mark-added.png", "句子标记：句内高亮、标签列表与标记计数")
    callout(doc, "小技巧", "点句子里的高亮词，会直接生成这个词的词汇卡片；再点「← 返回」回到原句。")

    # ---------------- 九
    h1(doc, "九、日历复习")
    p(doc, "「日历」标签页按日期查看当天学习情况：新增生词、收藏的句子、分析记录数量，"
           "支持前后翻页与「今天」快捷跳转。")
    figure(doc, "13-day.png", "日历复习：按天汇总生词、句子与记录")

    # ---------------- 十
    h1(doc, "十、导出")
    p(doc, "点顶栏「导出」选择格式，浏览器会下载文件，同时在 exports/ 目录保存一份；"
           "也可以命令行执行 python app.py export --format md。")
    figure(doc, "07-export.png", "导出菜单：五种格式")
    table(doc,
          ["格式", "文件名", "典型用途"],
          [["Markdown", "englearn-<时间戳>.md", "笔记软件归档、发布到博客"],
           ["JSON", "englearn-<时间戳>.json", "完整备份与迁移"],
           ["CSV", "englearn-<时间戳>.csv", "Excel 统计，兼容中文"],
           ["Anki", "englearn-<时间戳>.txt", "直接导入 Anki 制作卡片"],
           ["HTML", "englearn-<时间戳>.html", "单文件分享或打印"]],
          widths=[3.4, 6.2, 6.4])

    # ---------------- 十一
    h1(doc, "十一、常见问题")
    table(doc,
          ["现象", "原因", "处理"],
          [["双击后窗口一闪而过", "没找到 Python", "安装 Python 3.10+ 并勾选 PATH，或设 ENGL_PY"],
           ["浏览器白屏", "打开早于服务就绪", "刷新，或访问控制台里的地址"],
           ["端口打不开", "控制台被关闭", "重新启动，别关黑窗口"],
           ["未配置 API Key", "缺少密钥", "设置里填入并保存"],
           ["热键没反应", "组合键被占用", "看顶栏显示的实际热键，或换一组"],
           ["选中不弹窗", "安全软件拦截钩子", "看 logs/error.log，或加 --no-popup"],
           ["提示没有捕获到文本", "该程序不允许模拟复制", "改用剪贴板监听或直接粘贴"]],
          widths=[4.4, 4.6, 7.4])

    # ---------------- 十二
    h1(doc, "十二、附录")
    h2(doc, "12.1 目录结构")
    code(doc, """Englearn/
├─ start.cmd             Windows 一键启动
├─ run.sh                Linux / macOS 启动
├─ app.py                HTTP 服务 + 热键 + 悬浮窗
├─ core/                 配置、存储、模型调用、取词、导出
├─ web/                  前端页面（原生 HTML/CSS/JS）
├─ scripts/              测试、截图、文档生成
├─ docs/                 部署 / 发布 / 本教程
├─ data/app.db           本地数据库（自动生成）
└─ exports/              导出文件（自动生成）""")

    h2(doc, "12.2 常用命令")
    code(doc, """python app.py                     # 正常启动
python app.py --mock --no-open    # 演示模式，不自动开浏览器
python app.py --port 9000         # 指定端口
python app.py --no-popup          # 关闭悬浮窗，只保留页面
python app.py export --format md  # 命令行导出
python scripts/selftest.py        # 离线冒烟测试""")

    h2(doc, "12.3 数据与隐私")
    p(doc, "服务只监听 127.0.0.1，不对局域网开放；学习数据全部保存在 data/app.db。"
           "只有你主动分析时，该段文本才会发送给你自己配置的模型服务。"
           "config.json 中是明文 API Key，不要提交到代码仓库，也不要随截图外发。")

    return doc


def main() -> None:
    doc = build()
    doc.save(OUT)
    print("saved:", OUT, "figures:", _fig_no["n"])


if __name__ == "__main__":
    main()
